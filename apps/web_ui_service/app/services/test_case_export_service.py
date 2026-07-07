from __future__ import annotations

from copy import copy
from io import BytesIO
from pathlib import Path

try:
    from openpyxl import load_workbook  # type: ignore[import-untyped]
except Exception as exc:  # pragma: no cover - import guard for lightweight runtime images
    load_workbook = None  # type: ignore[assignment]
    _OPENPYXL_IMPORT_ERROR = exc
else:
    _OPENPYXL_IMPORT_ERROR = None

from app.models.test_case import TestCase
from app.services.test_case_data_service import render_test_steps_text

REPO_ROOT = Path(__file__).resolve().parents[4]
EXPORT_TEMPLATE_PATH = REPO_ROOT / "docs" / "用例模版示例.xlsx"
EXPORT_TEMPLATE_SHEET = "用例示例"
EXPORT_FILENAME = "test-cases-template.xlsx"
EXPORT_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

TEMPLATE_HEADERS = [
    "用例ID",
    "用例标题",
    "链路环节",
    "域/服务",
    "关联服务",
    "优先级",
    "测试类型",
    "触发入口",
    "注入点类型",
    "注入点目标",
    "注入参数",
    "前置数据SQL",
    "前置依赖状态",
    "测试步骤",
    "并发模型",
    "重试策略",
    "期望结果",
    "断言SQL",
    "事件断言",
    "指标断言",
    "回滚/清理脚本",
    "产物链接",
    "备注",
]


def _text(value: object, *, default: str = "") -> str:
    normalized = str(value or "").strip()
    return normalized or default


def _list_text(value: object) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]


def _comma_join(value: object) -> str:
    return ", ".join(_list_text(value))


def _newline_join(values: list[str]) -> str:
    return "\n".join(item for item in values if item)


def _with_na(value: object) -> str:
    return _text(value, default="N/A")


def _render_steps(case: TestCase) -> str:
    direct_text = _text(case.test_steps_text)
    if direct_text:
        return direct_text
    steps = case.test_steps if isinstance(case.test_steps, list) else []
    return render_test_steps_text(steps)


def _artifact_links(case: TestCase) -> str:
    links = _list_text(case.artifact_links)
    if links:
        return _newline_join(links)
    last_report_url = _text(case.last_report_url)
    if last_report_url:
        return last_report_url
    return ""


def build_template_row(case: TestCase) -> list[str]:
    return [
        _text(case.case_id),
        _text(case.name),
        _text(case.chain_stage),
        _text(case.sut_service),
        _comma_join(case.related_services),
        _text(case.priority, default="P2"),
        _comma_join(case.scenario_types),
        _text(case.trigger_entry),
        _with_na(case.fault_injection_type),
        _with_na(case.fault_injection_target),
        _with_na(case.fault_injection_params),
        _text(case.setup_sql),
        _text(case.precondition_state),
        _render_steps(case),
        _with_na(case.concurrency_model),
        _with_na(case.retry_policy),
        _text(case.expected_result),
        _text(case.assert_sql),
        _text(case.event_assertion),
        _text(case.metric_assertion),
        _text(case.cleanup_script),
        _artifact_links(case),
        _text(case.notes),
    ]


def _load_template():
    if load_workbook is None:
        message = "openpyxl is required for xlsx export. install dependency: pip install openpyxl>=3.1,<4.0"
        if _OPENPYXL_IMPORT_ERROR is not None:
            message = f"{message} (import error: {_OPENPYXL_IMPORT_ERROR})"
        raise RuntimeError(message)
    if not EXPORT_TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Excel export template not found: {EXPORT_TEMPLATE_PATH}")
    workbook = load_workbook(EXPORT_TEMPLATE_PATH)
    if EXPORT_TEMPLATE_SHEET not in workbook.sheetnames:
        raise ValueError(f"Excel export template sheet missing: {EXPORT_TEMPLATE_SHEET}")
    worksheet = workbook[EXPORT_TEMPLATE_SHEET]
    actual_headers = [worksheet.cell(row=1, column=index).value for index in range(1, len(TEMPLATE_HEADERS) + 1)]
    if actual_headers != TEMPLATE_HEADERS:
        raise ValueError("Excel export template headers do not match expected mapping")
    return workbook, worksheet


def build_export_xlsx(cases: list[TestCase]) -> bytes:
    workbook, worksheet = _load_template()
    style_source_row = 2 if worksheet.max_row >= 2 else 1
    row_styles = [
        copy(worksheet.cell(row=style_source_row, column=index)._style)
        for index in range(1, len(TEMPLATE_HEADERS) + 1)
    ]
    row_height = worksheet.row_dimensions[style_source_row].height

    if worksheet.max_row >= 2:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    for row_index, case in enumerate(cases, start=2):
        row_values = build_template_row(case)
        for column_index, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row=row_index, column=column_index)
            cell.value = value
            cell._style = copy(row_styles[column_index - 1])
        if row_height is not None:
            worksheet.row_dimensions[row_index].height = row_height

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
