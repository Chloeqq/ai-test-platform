from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

import app.models.test_case as test_case_model
from app.services import test_case_export_service


def _build_case(**overrides: object) -> test_case_model.TestCase:
    payload: dict[str, object] = {
        "case_id": "atp-web-ret-query-sm-ai-0001",
        "project_code": "atp",
        "client": "web",
        "page_code": "ret",
        "page_name": "退货",
        "module_code": "query",
        "module_name": "查询",
        "case_type": "sm",
        "source": "ai",
        "name": "退货查询冒烟校验",
        "product_line": "退货",
        "module": "查询",
        "chain_stage": "退货/查询",
        "sut_service": "return-service",
        "related_services": ["auth-service", "search-service"],
        "priority": "P1",
        "test_type": "ui",
        "scenario_types": ["并发", "幂等"],
        "trigger_entry": "API",
        "fault_injection_type": "",
        "fault_injection_target": "",
        "fault_injection_params": "",
        "setup_sql": "SELECT 1;",
        "precondition_state": "用户已登录",
        "test_steps": [],
        "test_steps_text": "Step1 打开退货页\nStep2 输入服务单号\nStep3 点击查询",
        "concurrency_model": "",
        "retry_policy": "",
        "expected_result": "返回退货单详情",
        "assert_sql": "SELECT count(*) FROM returns;",
        "event_assertion": "OrderRefundQueried 触发 1 次",
        "metric_assertion": "查询失败率 = 0",
        "cleanup_script": "DELETE FROM returns WHERE id='demo';",
        "artifact_links": ["https://report.local/trace/1", "https://report.local/log/1"],
        "notes": "需关注老数据兼容",
        "tags": ["smoke"],
        "markers": ["p1"],
        "creator": "qa",
        "assignee": "qa-owner",
        "pytest_path": "tests/test_return_query.py",
        "status": "active",
        "automation_status": "automated",
        "created_source": "ai",
        "source_ref": "REQ-1",
        "script_code": "def test_case():\n    assert True\n",
        "data_config": {},
        "last_execution_result": "passed",
        "last_report_url": "https://report.local/latest",
    }
    payload.update(overrides)
    return test_case_model.TestCase(**payload)


def test_build_export_xlsx_uses_template_headers_and_case_id() -> None:
    workbook_bytes = test_case_export_service.build_export_xlsx([_build_case()])

    workbook = load_workbook(BytesIO(workbook_bytes))
    worksheet = workbook[test_case_export_service.EXPORT_TEMPLATE_SHEET]
    headers = [worksheet.cell(row=1, column=index).value for index in range(1, 24)]

    assert worksheet.freeze_panes == "A2"
    assert headers == test_case_export_service.TEMPLATE_HEADERS
    assert worksheet.cell(row=2, column=1).value == "atp-web-ret-query-sm-ai-0001"
    assert worksheet.cell(row=2, column=2).value == "退货查询冒烟校验"
    assert worksheet.cell(row=2, column=5).value == "auth-service, search-service"
    assert worksheet.cell(row=2, column=7).value == "并发, 幂等"


def test_build_export_xlsx_applies_template_fallbacks() -> None:
    workbook_bytes = test_case_export_service.build_export_xlsx(
        [
            _build_case(
                artifact_links=[],
                last_report_url="https://report.local/fallback",
                test_steps_text="",
                test_steps=[
                    {"description": "打开页面"},
                    {"description": "点击查询", "expected": "结果出现"},
                ],
            )
        ]
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    worksheet = workbook[test_case_export_service.EXPORT_TEMPLATE_SHEET]

    assert worksheet.cell(row=2, column=9).value == "N/A"
    assert worksheet.cell(row=2, column=10).value == "N/A"
    assert worksheet.cell(row=2, column=11).value == "N/A"
    assert worksheet.cell(row=2, column=15).value == "N/A"
    assert worksheet.cell(row=2, column=16).value == "N/A"
    assert worksheet.cell(row=2, column=14).value == "Step1 打开页面\nStep2 点击查询 | 结果出现"
    assert worksheet.cell(row=2, column=22).value == "https://report.local/fallback"
